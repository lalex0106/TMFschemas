#!/usr/bin/env python3
"""基于企业 Schema 输出生成翻译工作簿并回写 YAML 文件的辅助脚本。

使用场景：
1. 在流水线完成 `dist/json` 目录的 Schema 生成后，执行 `extract` 子命令，将需要翻译
   的模型名称、描述及属性说明汇总到单一 Excel 文件中，便于翻译团队批量处理。
2. 翻译人员通过机器翻译或人工校对更新 Excel 中的中文列后，执行 `render` 子命令，可
   自动生成符合流水线约定的 YAML 翻译文件目录结构，无需手动维护零散的 YAML/CSV。

脚本依赖 `openpyxl` 与 `PyYAML`，如未预装可通过 `pip install openpyxl pyyaml` 获取。
"""
from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, MutableMapping, Optional, Sequence, Tuple

try:  # pragma: no cover - 缺失依赖时提示
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]

try:  # pragma: no cover - 缺失依赖时提示
    import yaml
except ImportError:  # pragma: no cover
    yaml = None  # type: ignore[assignment]


DEFAULT_SCHEMA_SHEET = "Schemas"
DEFAULT_PROPERTY_SHEET = "Properties"
SCHEMA_HEADERS: Tuple[str, ...] = (
    "Domain",
    "Schema",
    "文件路径",
    "英文名称",
    "英文描述",
    "中文名称",
    "中文描述",
)
PROPERTY_HEADERS: Tuple[str, ...] = (
    "Domain",
    "Schema",
    "属性路径",
    "英文名称",
    "英文描述",
    "中文名称",
    "中文描述",
)
GLOBAL_PROPERTY_KEY = "__GLOBAL_PROPERTIES__"


@dataclass
class ExtractedSchema:
    domain: str
    schema_file: str
    schema_name: str
    title_en: str
    description_en: str
    title_locale: str
    description_locale: str


@dataclass
class ExtractedProperty:
    domain: str
    schema_file: str
    schema_name: str
    property_path: str
    title_en: str
    description_en: str
    title_locale: str
    description_locale: str


def _require_openpyxl(context: str) -> None:
    if Workbook is None or load_workbook is None:  # pragma: no cover - 仅在缺失依赖时触发
        raise SystemExit(
            f"{context} 需要 openpyxl 支持，请先执行 `pip install openpyxl` 或在隔离环境中预装该依赖。"
        )


def _require_yaml(context: str) -> None:
    if yaml is None:  # pragma: no cover - 仅在缺失依赖时触发
        raise SystemExit(
            f"{context} 需要 PyYAML 支持，请先执行 `pip install pyyaml` 或在隔离环境中预装该依赖。"
        )


def _load_json(path: Path) -> Optional[Mapping[str, object]]:
    try:
        with path.open("r", encoding="utf-8") as fh:
            data = json.load(fh)
    except json.JSONDecodeError as exc:
        print(f"⚠️  解析失败 {path}: {exc}")
        return None
    return data if isinstance(data, Mapping) else None


def _normalize_str(value: object) -> str:
    if isinstance(value, str):
        return value.strip()
    return ""


def _extract_locale_value(block: Mapping[str, object], field: str, locale: str) -> str:
    target = block.get(field)
    if isinstance(target, Mapping):
        text = target.get(locale)
        if isinstance(text, str):
            return text.strip()
    return ""


def _iter_property_nodes(
    node: Mapping[str, object],
    path: Tuple[str, ...] = (),
) -> Iterable[Tuple[Tuple[str, ...], Mapping[str, object]]]:
    properties = node.get("properties")
    if isinstance(properties, Mapping):
        for prop_name, prop_value in properties.items():
            if not isinstance(prop_value, Mapping):
                continue
            new_path = path + (str(prop_name),)
            yield new_path, prop_value
            yield from _iter_property_nodes(prop_value, new_path)

    items = node.get("items")
    if isinstance(items, Mapping):
        yield from _iter_property_nodes(items, path)

    additional = node.get("additionalProperties")
    if isinstance(additional, Mapping):
        yield from _iter_property_nodes(additional, path)

    for key in ("allOf", "anyOf", "oneOf"):
        collection = node.get(key)
        if isinstance(collection, list):
            for index, item in enumerate(collection):
                if isinstance(item, Mapping):
                    yield from _iter_property_nodes(item, path)


def _format_property_path(path: Tuple[str, ...]) -> str:
    return ".".join(path)


def _collect_schema_entries(
    schema_dir: Path,
    locale: str,
) -> Tuple[List[ExtractedSchema], List[ExtractedProperty]]:
    schemas: List[ExtractedSchema] = []
    properties: List[ExtractedProperty] = []
    seen_properties: set[Tuple[str, Tuple[str, ...]]] = set()

    for file in sorted(schema_dir.rglob("*.schema.json")):
        payload = _load_json(file)
        if not payload:
            continue
        definitions = payload.get("definitions")
        if not isinstance(definitions, Mapping):
            continue

        try:
            relative = file.relative_to(schema_dir)
            parts = relative.parts
            schema_file = str(relative)
            domain = parts[0] if len(parts) > 1 else ""
        except ValueError:
            schema_file = str(file)
            domain = file.parent.name if file.parent != schema_dir else ""

        for schema_name, definition in definitions.items():
            if not isinstance(definition, Mapping):
                continue
            title_en = _normalize_str(definition.get("title"))
            description_en = _normalize_str(definition.get("description"))
            x_i18n = definition.get("x-i18n")
            locale_title = ""
            locale_desc = ""
            if isinstance(x_i18n, Mapping):
                locale_title = _extract_locale_value(x_i18n, "title", locale)
                locale_desc = _extract_locale_value(x_i18n, "description", locale)

            schemas.append(
                ExtractedSchema(
                    domain=domain,
                    schema_file=schema_file,
                    schema_name=str(schema_name),
                    title_en=title_en,
                    description_en=description_en,
                    title_locale=locale_title,
                    description_locale=locale_desc,
                )
            )

            for path_tuple, prop_node in _iter_property_nodes(definition, ()):  # type: ignore[arg-type]
                key = (schema_name, path_tuple)
                if key in seen_properties:
                    continue
                seen_properties.add(key)

                title_en_prop = _normalize_str(prop_node.get("title"))
                description_en_prop = _normalize_str(prop_node.get("description"))

                locale_title_prop = ""
                locale_desc_prop = ""
                prop_i18n = prop_node.get("x-i18n")
                if isinstance(prop_i18n, Mapping):
                    locale_title_prop = _extract_locale_value(prop_i18n, "title", locale)
                    locale_desc_prop = _extract_locale_value(prop_i18n, "description", locale)

                properties.append(
                    ExtractedProperty(
                        domain=domain,
                        schema_file=schema_file,
                        schema_name=str(schema_name),
                        property_path=_format_property_path(path_tuple),
                        title_en=title_en_prop,
                        description_en=description_en_prop,
                        title_locale=locale_title_prop,
                        description_locale=locale_desc_prop,
                    )
                )

    return schemas, properties


def _write_workbook(
    workbook_path: Path,
    schemas: Sequence[ExtractedSchema],
    properties: Sequence[ExtractedProperty],
) -> None:
    _require_openpyxl("导出翻译工作簿")
    wb = Workbook()
    if wb.worksheets:
        ws = wb.active
        ws.title = DEFAULT_SCHEMA_SHEET
    else:  # pragma: no cover - openpyxl 始终会创建一个默认工作表
        ws = wb.create_sheet(DEFAULT_SCHEMA_SHEET)

    ws.append(list(SCHEMA_HEADERS))
    for row in schemas:
        ws.append(
            [
                row.domain,
                row.schema_name,
                row.schema_file,
                row.title_en,
                row.description_en,
                row.title_locale,
                row.description_locale,
            ]
        )

    prop_sheet = wb.create_sheet(DEFAULT_PROPERTY_SHEET)
    prop_sheet.append(list(PROPERTY_HEADERS))
    for row in properties:
        prop_sheet.append(
            [
                row.domain,
                row.schema_name,
                row.property_path,
                row.title_en,
                row.description_en,
                row.title_locale,
                row.description_locale,
            ]
        )

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(workbook_path)
    print(f"✅ 已生成翻译工作簿：{workbook_path}")


def command_extract(args: argparse.Namespace) -> None:
    schema_dir = Path(args.schema_dir)
    if not schema_dir.exists():
        raise SystemExit(f"未找到 schema 目录：{schema_dir}")
    locale = args.locale

    schemas, properties = _collect_schema_entries(schema_dir, locale)
    if not schemas:
        print("⚠️ 未在目标目录中发现 definitions，请确认已执行构建流水线。")

    workbook_path = Path(args.workbook)
    _write_workbook(workbook_path, schemas, properties)


def _sanitize_header(value: Optional[object]) -> str:
    if isinstance(value, str):
        return value.strip().lstrip("\ufeff")
    return ""


def _map_headers(headers: Sequence[str]) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for idx, name in enumerate(headers):
        lowered = name.lower().replace(" ", "")
        mapping[lowered] = idx
    return mapping


def _locate_column(header_map: Mapping[str, int], candidates: Iterable[str]) -> Optional[int]:
    normalized_candidates = [c.lower().replace(" ", "") for c in candidates]
    for candidate in normalized_candidates:
        if candidate in header_map:
            return header_map[candidate]
    return None


def command_render(args: argparse.Namespace) -> None:
    _require_openpyxl("加载翻译工作簿")
    _require_yaml("生成 YAML 翻译文件")

    workbook_path = Path(args.workbook)
    if not workbook_path.exists():
        raise SystemExit(f"未找到工作簿：{workbook_path}")

    wb = load_workbook(workbook_path)
    schema_sheet = wb[args.schema_sheet] if args.schema_sheet in wb.sheetnames else wb[DEFAULT_SCHEMA_SHEET]
    property_sheet = (
        wb[args.property_sheet]
        if args.property_sheet in wb.sheetnames
        else wb[DEFAULT_PROPERTY_SHEET]
    )

    header_row = next(schema_sheet.iter_rows(min_row=1, max_row=1), None)
    if header_row is None:
        raise SystemExit("Schema 工作表为空，请至少保留表头行。")
    schema_headers = [_sanitize_header(cell.value) for cell in header_row]

    property_header_row = next(property_sheet.iter_rows(min_row=1, max_row=1), None)
    if property_header_row is None:
        raise SystemExit("属性工作表为空，请至少保留表头行。")
    property_headers = [_sanitize_header(cell.value) for cell in property_header_row]

    schema_map = _map_headers(schema_headers)
    property_map = _map_headers(property_headers)

    schema_col = _locate_column(schema_map, ["schema", "模型", "name"])
    schema_title_col = _locate_column(schema_map, ["中文名称", "titlezh", "title_zh", "translatedtitle"])
    schema_desc_col = _locate_column(schema_map, ["中文描述", "新描述", "descriptionzh", "translateddescription"])

    if schema_col is None:
        raise SystemExit("工作簿中缺少 Schema 列，请确认表头是否包含 Schema/模型 名称。")

    schema_payloads: Dict[str, MutableMapping[str, object]] = {}

    for row in schema_sheet.iter_rows(min_row=2):
        schema_name = _sanitize_header(row[schema_col].value)
        if not schema_name:
            continue
        title_text = _sanitize_header(row[schema_title_col].value) if schema_title_col is not None else ""
        desc_text = _sanitize_header(row[schema_desc_col].value) if schema_desc_col is not None else ""

        if not title_text and not desc_text:
            continue

        payload = schema_payloads.setdefault(schema_name, {"schema": schema_name, "translations": {}})
        translations = payload.setdefault("translations", {})
        if not isinstance(translations, MutableMapping):
            translations = dict(translations)
            payload["translations"] = translations
        if title_text:
            translations["title"] = title_text
        if desc_text:
            translations["description"] = desc_text

    prop_schema_col = _locate_column(property_map, ["schema", "模型"])
    prop_path_col = _locate_column(property_map, ["属性路径", "property", "path"])
    prop_title_col = _locate_column(property_map, ["中文名称", "titlezh", "title_zh", "translatedtitle"])
    prop_desc_col = _locate_column(property_map, ["中文描述", "新描述", "descriptionzh", "translateddescription"])

    if prop_path_col is None:
        raise SystemExit("属性工作表缺少属性路径列，请确认表头是否包含 属性路径/Property/Path。")

    for row in property_sheet.iter_rows(min_row=2):
        schema_name = _sanitize_header(row[prop_schema_col].value) if prop_schema_col is not None else ""
        property_path = _sanitize_header(row[prop_path_col].value)
        if not property_path:
            continue
        title_text = _sanitize_header(row[prop_title_col].value) if prop_title_col is not None else ""
        desc_text = _sanitize_header(row[prop_desc_col].value) if prop_desc_col is not None else ""
        if not title_text and not desc_text:
            continue

        target_key = schema_name or GLOBAL_PROPERTY_KEY
        payload = schema_payloads.setdefault(target_key, {"schema": target_key, "translations": {}})
        translations = payload.setdefault("translations", {})
        if not isinstance(translations, MutableMapping):
            translations = dict(translations)
            payload["translations"] = translations

        path_parts = [segment for segment in property_path.split(".") if segment]
        current = translations.setdefault("properties", {})
        if not isinstance(current, MutableMapping):
            current = dict(current)
            translations["properties"] = current

        node = current
        for part in path_parts[:-1]:
            if part == "[]":
                items = node.get("items")
                if isinstance(items, MutableMapping):
                    node = items
                elif isinstance(items, Mapping):
                    new_items: MutableMapping[str, object] = dict(items)
                    node["items"] = new_items
                    node = new_items
                else:
                    new_items = {}
                    node["items"] = new_items
                    node = new_items
                continue
            if part == "{}":
                additional = node.get("additionalProperties")
                if isinstance(additional, MutableMapping):
                    node = additional
                elif isinstance(additional, Mapping):
                    new_additional: MutableMapping[str, object] = dict(additional)
                    node["additionalProperties"] = new_additional
                    node = new_additional
                else:
                    new_additional = {}
                    node["additionalProperties"] = new_additional
                    node = new_additional
                continue
            props = node.get("properties")
            if isinstance(props, MutableMapping):
                properties_map = props
            elif isinstance(props, Mapping):
                properties_map = dict(props)
                node["properties"] = properties_map
            else:
                properties_map = {}
                node["properties"] = properties_map
            nested = properties_map.get(part)
            if isinstance(nested, MutableMapping):
                node = nested
            elif isinstance(nested, Mapping):
                new_nested: MutableMapping[str, object] = dict(nested)
                properties_map[part] = new_nested
                node = new_nested
            else:
                new_nested = {}
                properties_map[part] = new_nested
                node = new_nested

        leaf_key = path_parts[-1] if path_parts else ""
        if leaf_key in {"[]", "{}"}:
            leaf_container = node.setdefault(
                "items" if leaf_key == "[]" else "additionalProperties", {}
            )
            if not isinstance(leaf_container, MutableMapping):
                leaf_container = dict(leaf_container) if isinstance(leaf_container, Mapping) else {}
                node["items" if leaf_key == "[]" else "additionalProperties"] = leaf_container
            leaf = leaf_container
        else:
            props = node.setdefault("properties", {}) if leaf_key else node
            if leaf_key:
                if not isinstance(props, MutableMapping):
                    props = dict(props) if isinstance(props, Mapping) else {}
                    node["properties"] = props
                leaf = props.setdefault(leaf_key, {})
                if not isinstance(leaf, MutableMapping):
                    leaf = dict(leaf) if isinstance(leaf, Mapping) else {}
                    props[leaf_key] = leaf
            else:
                leaf = props  # type: ignore[assignment]

        if isinstance(leaf, MutableMapping):
            if title_text:
                leaf["title"] = title_text
            if desc_text:
                leaf["description"] = desc_text

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    for schema_name, payload in sorted(schema_payloads.items()):
        output_path = output_dir / f"{schema_name}.yaml"
        with output_path.open("w", encoding="utf-8") as fh:
            yaml.safe_dump(payload, fh, allow_unicode=True, sort_keys=False)
        print(f"✅ 已生成翻译 YAML：{output_path}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="翻译工作簿导入导出工具")
    subparsers = parser.add_subparsers(dest="command", required=True)

    extract_parser = subparsers.add_parser("extract", help="从 dist/json 导出 Excel 工作簿")
    extract_parser.add_argument(
        "--schema-dir",
        default="dist/json",
        help="已构建的 schema 根目录，默认 dist/json",
    )
    extract_parser.add_argument(
        "--workbook",
        default="overrides/i18n/workbooks/translations.xlsx",
        help="输出的 Excel 路径，默认 overrides/i18n/workbooks/translations.xlsx",
    )
    extract_parser.add_argument(
        "--locale",
        default="zh-CN",
        help="读取 `x-i18n` 时优先填充的语种编码，默认 zh-CN",
    )
    extract_parser.set_defaults(func=command_extract)

    render_parser = subparsers.add_parser("render", help="从 Excel 生成 YAML 翻译文件")
    render_parser.add_argument(
        "--workbook",
        default="overrides/i18n/workbooks/translations.xlsx",
        help="翻译后的 Excel 文件路径",
    )
    render_parser.add_argument(
        "--output-dir",
        default="overrides/i18n/zh-CN/yaml",
        help="生成 YAML 文件的输出目录，默认 overrides/i18n/zh-CN/yaml",
    )
    render_parser.add_argument(
        "--schema-sheet",
        default=DEFAULT_SCHEMA_SHEET,
        help="模型翻译所在的工作表名称，默认 Schemas",
    )
    render_parser.add_argument(
        "--property-sheet",
        default=DEFAULT_PROPERTY_SHEET,
        help="属性翻译所在的工作表名称，默认 Properties",
    )
    render_parser.set_defaults(func=command_render)

    return parser


def main(argv: Optional[Sequence[str]] = None) -> None:
    parser = build_parser()
    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":  # pragma: no cover - 脚本入口
    main()
